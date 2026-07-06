"""
core/excel_report.py
--------------------
Generates a professional Excel calibration report using openpyxl,
mirroring the column layout of the PDF report and the UI TX/RX pages.

TX columns:  BAND | BLK | FREQ | MOD | BW | ANT | ORIGIN dBm | DUT dBm |
             DELTA | CORRECTION | TARGET | LIMITS | STATUS
RX columns:  BAND | BLK | FREQ | MCS | BW | ANT | ORIGIN RSSI | DUT RSSI |
             Δ RSSI | ORIG PER % | DUT PER % | STATUS

Install:  pip install openpyxl

Usage
-----
    from core.excel_report import ExcelReport
    from core import CalibrationEngine

    report = CalibrationEngine.compare(origin, dut)
    ExcelReport.generate(report, "output/calibration_report.xlsx",
                         operator="Ahmed", session_id=42)
"""

from pathlib import Path
from datetime import datetime
from typing import Optional

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

# ── Colour palette (matches PDF + dark-UI theme) ───────────────────────────────
DARK        = "0D1117"
ACCENT      = "E8A020"
PASS_CLR    = "2ECC71"
FAIL_CLR    = "E74C3C"
WARN_CLR    = "E8A020"
C_5G        = "3B8FD4"
C_2G        = "27AE60"
C_6G        = "E67E22"
HEADER_BG   = "1C2030"
ROW_WHITE   = "FFFFFF"
ROW_ALT     = "F4F6FA"
BORDER_CLR  = "C8CDD8"
WHITE       = "FFFFFF"
GRAY        = "6C7A8A"
PASS_BG     = "0D2A1A"
FAIL_BG     = "2A0D0D"
WARN_ROW1   = "FFF8EC"
WARN_ROW2   = "FFF3D6"

BAND_COLORS = {"5 GHz": C_5G, "2.4 GHz": C_2G, "6 GHz": C_6G}

# ── Low-level helpers ──────────────────────────────────────────────────────────

def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _font(bold=False, color=DARK, size=9, name="Arial", italic=False):
    return Font(bold=bold, color=color, size=size, name=name, italic=italic)


def _align(h="center", v="center", wrap=False):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)


def _thin_border(color=BORDER_CLR):
    s = Side(border_style="thin", color=color)
    return Border(left=s, right=s, top=s, bottom=s)


def _set_col_widths(ws, widths: list):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _write_header_row(ws, row_num: int, labels: list, bg: str = HEADER_BG):
    """Write a styled dark-background header row."""
    for col, label in enumerate(labels, 1):
        c = ws.cell(row=row_num, column=col, value=label)
        c.font = _font(bold=True, color=WHITE, size=8)
        c.fill = _fill(bg)
        c.alignment = _align()
        c.border = _thin_border()
    ws.row_dimensions[row_num].height = 20


def _write_data_cell(ws, row_num: int, col: int, value,
                     color: str = DARK, bold: bool = False,
                     bg: str = ROW_WHITE, mono: bool = True,
                     border_color: str = BORDER_CLR):
    c = ws.cell(row=row_num, column=col, value=value)
    c.font = Font(
        bold=bold, color=color, size=8,
        name="Courier New" if (mono and not bold) else "Arial"
    )
    c.fill = _fill(bg)
    c.alignment = _align()
    c.border = _thin_border(border_color)
    return c


# ── Public API ─────────────────────────────────────────────────────────────────

class ExcelReport:

    @staticmethod
    def generate(
        report,
        output_path: str,
        operator: str = "",
        session_id: Optional[int] = None,
    ) -> str:
        """
        Generate a multi-sheet Excel calibration report from a CalibrationReport.
        Returns the output path.
        """
        if not OPENPYXL_AVAILABLE:
            raise ImportError(
                "openpyxl is required for Excel export.\n"
                "Install it with:  pip install openpyxl"
            )

        output_path = str(output_path)
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)

        wb = Workbook()

        ws_sum = wb.active
        ws_sum.title = "Summary"
        _build_summary_sheet(ws_sum, report, operator, session_id)

        ws_tx = wb.create_sheet("TX Results")
        _build_tx_sheet(ws_tx, report)

        ws_rx = wb.create_sheet("RX Results")
        _build_rx_sheet(ws_rx, report)

        ws_cor = wb.create_sheet("Corrections")
        _build_corrections_sheet(ws_cor, report)

        wb.save(output_path)
        return output_path


# ── Sheet builders ─────────────────────────────────────────────────────────────

def _build_summary_sheet(ws, report, operator: str, session_id):
    ws.sheet_view.showGridLines = False

    # ── Title ──────────────────────────────────────────────────────────────────
    ws.row_dimensions[1].height = 32
    ws.merge_cells("A1:L1")
    t = ws.cell(row=1, column=1, value="WiFi Calibration Report")
    t.font = Font(bold=True, size=18, color=DARK, name="Arial")
    t.alignment = _align(h="left", v="center")

    # ── Metadata block ─────────────────────────────────────────────────────────
    now = datetime.now().strftime("%Y-%m-%d  %H:%M")
    sid = f"Session #{session_id}" if session_id else "—"
    meta = [
        ("Product",        report.product_name),
        ("DUT Serial",     report.dut_serial),
        ("Origin Serial",  report.origin_serial),
        ("Operator",       operator or "—"),
        ("Date",           now),
        ("Session",        sid),
    ]
    for i, (lbl, val) in enumerate(meta, 2):
        lc = ws.cell(row=i, column=1, value=f"{lbl}:")
        lc.font = _font(bold=True, color=GRAY, size=9)
        lc.alignment = _align(h="left")
        vc = ws.cell(row=i, column=2, value=val)
        vc.font = _font(color=DARK, size=9)
        vc.alignment = _align(h="left")

    # ── Amber separator ────────────────────────────────────────────────────────
    sep = len(meta) + 3
    ws.row_dimensions[sep].height = 4
    for col in range(1, 13):
        ws.cell(row=sep, column=col).fill = _fill(ACCENT)

    # ── Result banner ──────────────────────────────────────────────────────────
    banner_row = sep + 2
    ws.row_dimensions[banner_row].height = 30
    ws.merge_cells(f"A{banner_row}:L{banner_row}")
    if report.overall_pass:
        banner_text = "OVERALL RESULT:  PASS  — No corrections required"
        banner_bg, banner_fg = PASS_BG, PASS_CLR
    else:
        banner_text = "OVERALL RESULT:  FAIL / CORRECTION NEEDED"
        banner_bg, banner_fg = FAIL_BG, FAIL_CLR
    bc = ws.cell(row=banner_row, column=1, value=banner_text)
    bc.font = Font(bold=True, size=13, color=banner_fg, name="Arial")
    bc.fill = _fill(banner_bg)
    bc.alignment = _align(h="center")

    # ── 8-metric tile row ──────────────────────────────────────────────────────
    lbl_row = banner_row + 2
    val_row = lbl_row + 1
    ws.row_dimensions[lbl_row].height = 14
    ws.row_dimensions[val_row].height = 24

    avg = f"{report.tx_avg_delta:+.3f}" if report.tx_avg_delta is not None else "—"
    mx  = f"{report.tx_max_delta:.3f}"  if report.tx_max_delta  is not None else "—"
    tiles = [
        ("TX BLOCKS",    str(len(report.tx_results)), DARK),
        ("TX PASS",      str(report.tx_pass_count),   PASS_CLR),
        ("TX FAIL",      str(report.tx_fail_count),   FAIL_CLR if report.tx_fail_count else DARK),
        ("CORRECTIONS",  str(report.tx_needs_correction_count),
                         WARN_CLR if report.tx_needs_correction_count else DARK),
        ("RX PASS",      str(report.rx_pass_count),   PASS_CLR),
        ("RX FAIL",      str(report.rx_fail_count),   FAIL_CLR if report.rx_fail_count else DARK),
        ("AVG Δ",        avg,                          DARK),
        ("MAX |Δ|",      mx,                           DARK),
    ]
    for col, (lbl, val, clr) in enumerate(tiles, 1):
        lc = ws.cell(row=lbl_row, column=col, value=lbl)
        lc.font = Font(size=7, color=GRAY, name="Arial")
        lc.fill = _fill(ROW_ALT)
        lc.alignment = _align()
        lc.border = _thin_border()

        vc = ws.cell(row=val_row, column=col, value=val)
        vc.font = Font(bold=True, size=13, color=clr, name="Courier New")
        vc.fill = _fill(ROW_ALT)
        vc.alignment = _align()
        vc.border = _thin_border()

    # ── Shortcuts to other sheets ──────────────────────────────────────────────
    link_row = val_row + 3
    links = [
        ("TX Results →",   "TX Results"),
        ("RX Results →",   "RX Results"),
        ("Corrections →",  "Corrections"),
    ]
    for col, (label, sheet) in enumerate(links, 1):
        c = ws.cell(row=link_row, column=col, value=label)
        c.font = Font(bold=True, size=9, color=C_5G, name="Arial", underline="single")
        c.hyperlink = f"#{sheet}!A1"
        c.alignment = _align(h="left")

    # ── Footer ─────────────────────────────────────────────────────────────────
    foot_row = link_row + 2
    ws.merge_cells(f"A{foot_row}:L{foot_row}")
    ft = ws.cell(
        row=foot_row, column=1,
        value=f"Generated by WiFi Calibration Tool — Sagemcom / EDERSON  |  "
              f"{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}  |  Operator: {operator or 'N/A'}"
    )
    ft.font = Font(size=7, color=GRAY, italic=True, name="Arial")
    ft.alignment = _align(h="center")

    _set_col_widths(ws, [15, 28, 12, 12, 12, 12, 12, 12, 12, 12, 12, 12])


def _build_tx_sheet(ws, report):
    """
    TX Calibration Results — 13 columns matching the PDF and UI exactly:
    BAND | BLK | FREQ (MHz) | MOD | BW | ANT | ORIGIN dBm | DUT dBm |
    DELTA | CORRECTION | TARGET | LIMITS | STATUS
    """
    ws.sheet_view.showGridLines = False

    # Sheet title
    ws.merge_cells("A1:M1")
    ws.row_dimensions[1].height = 24
    t = ws.cell(row=1, column=1, value="TX Calibration Results")
    t.font = Font(bold=True, size=12, color=DARK, name="Arial")
    t.alignment = _align(h="left", v="center")
    t.fill = _fill(ROW_ALT)

    headers = [
        "BAND", "BLK", "FREQ (MHz)", "MOD", "BW", "ANT",
        "ORIGIN dBm", "DUT dBm", "DELTA", "CORRECTION", "TARGET", "LIMITS", "STATUS",
    ]
    _write_header_row(ws, 2, headers)

    for i, r in enumerate(report.tx_results):
        row = i + 3
        bg = ROW_ALT if (i % 2 == 1) else ROW_WHITE
        ws.row_dimensions[row].height = 17

        band_val = r.band.value if r.band else "?"
        band_clr = BAND_COLORS.get(band_val, GRAY)

        orig    = f"{r.origin_measured_dbm:.2f}" if r.origin_measured_dbm is not None else "—"
        dut     = f"{r.dut_measured_dbm:.2f}"    if r.dut_measured_dbm    is not None else "—"
        delta   = f"{r.delta_dbm:+.3f}"          if r.delta_dbm           is not None else "—"
        corr    = f"{r.correction_dbm:+.3f}"     if r.correction_dbm      is not None else "—"
        target  = f"{r.tx_target_dbm:.1f}"       if r.tx_target_dbm       is not None else "—"
        limits  = r.limits_str

        s = r.status.value
        if "PASS" in s or s == "OK":
            sc = PASS_CLR
        elif "FAIL" in s:
            sc = FAIL_CLR
        elif "CORR" in s:
            sc = WARN_CLR
        else:
            sc = GRAY

        # Delta background: amber tint if correction needed
        delta_bg = "FFF3D6" if r.needs_correction else bg
        corr_bg  = "FFF3D6" if r.needs_correction else bg

        row_data = [
            # (value, color, bold, bg_override, mono)
            (band_val,          band_clr, True,  bg,       False),
            (r.block_number,    DARK,     False, bg,       True),
            (r.freq_mhz,        DARK,     False, bg,       True),
            (r.modulation,      DARK,     False, bg,       True),
            (r.bandwidth,       DARK,     False, bg,       True),
            (r.antenna,         DARK,     False, bg,       True),
            (orig,              DARK,     False, bg,       True),
            (dut,               DARK,     False, bg,       True),
            (delta,             WARN_CLR if r.needs_correction else DARK,
                                          False, delta_bg, True),
            (corr,              WARN_CLR if r.needs_correction else DARK,
                                          True,  corr_bg,  False),
            (target,            GRAY,     False, bg,       True),
            (limits,            GRAY,     False, bg,       True),
            (s,                 sc,       True,  bg,       False),
        ]
        for col, (val, clr, bold, cell_bg, mono) in enumerate(row_data, 1):
            _write_data_cell(ws, row, col, val,
                             color=clr, bold=bold, bg=cell_bg, mono=mono)

    last_data_row = len(report.tx_results) + 2
    ws.auto_filter.ref = f"A2:M{last_data_row}"
    ws.freeze_panes = "A3"
    # col widths:  Band  Blk  Freq   Mod    BW     Ant    Orig   DUT   Delta  Corr   Target Limits Status
    _set_col_widths(ws, [10,   5,    10,    10,    8,     7,     13,   13,    11,    13,    9,     14,    15])


def _build_rx_sheet(ws, report):
    """
    RX Comparison Results — 12 columns matching the PDF and UI exactly:
    BAND | BLK | FREQ (MHz) | MCS | BW | ANT | ORIGIN RSSI | DUT RSSI |
    Δ RSSI | ORIG PER % | DUT PER % | STATUS
    """
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:L1")
    ws.row_dimensions[1].height = 24
    t = ws.cell(row=1, column=1, value="RX Comparison Results")
    t.font = Font(bold=True, size=12, color=DARK, name="Arial")
    t.alignment = _align(h="left", v="center")
    t.fill = _fill(ROW_ALT)

    headers = [
        "BAND", "BLK", "FREQ (MHz)", "MCS", "BW", "ANT",
        "ORIGIN RSSI", "DUT RSSI", "Δ RSSI",
        "ORIG PER %", "DUT PER %", "STATUS",
    ]
    _write_header_row(ws, 2, headers)

    for i, r in enumerate(report.rx_results):
        row = i + 3
        bg = ROW_ALT if (i % 2 == 1) else ROW_WHITE
        ws.row_dimensions[row].height = 17

        band_val = r.band.value if r.band else "?"
        band_clr = BAND_COLORS.get(band_val, GRAY)

        orig_rssi = f"{r.origin_rssi:.1f}" if (r.origin_rssi is not None and r.origin_rssi != -999) else "—"
        dut_rssi  = f"{r.dut_rssi:.1f}"    if (r.dut_rssi   is not None and r.dut_rssi   != -999) else "—"
        delta_r   = f"{r.rssi_delta:+.1f}" if r.rssi_delta  is not None else "—"
        orig_per  = f"{r.origin_per:.2f}"  if r.origin_per  is not None else "—"
        dut_per   = f"{r.dut_per:.2f}"     if r.dut_per     is not None else "—"

        s = r.status.value
        if "PASS" in s:
            sc = PASS_CLR
        elif "FAIL" in s:
            sc = FAIL_CLR
        elif s == "INVALID":
            sc = GRAY
        else:
            sc = GRAY

        # Dim INVALID rows slightly
        text_clr = GRAY if s == "INVALID" else DARK

        row_data = [
            (band_val,   band_clr,  True,  False),
            (r.block_number, text_clr, False, True),
            (r.freq_mhz, text_clr,  False, True),
            (r.mcs,      text_clr,  False, True),
            (r.bandwidth, text_clr, False, True),
            (r.antenna_label, text_clr, False, True),
            (orig_rssi,  text_clr,  False, True),
            (dut_rssi,   text_clr,  False, True),
            (delta_r,    text_clr,  False, True),
            (orig_per,   text_clr,  False, True),
            (dut_per,    text_clr,  False, True),
            (s,          sc,        True,  False),
        ]
        for col, (val, clr, bold, mono) in enumerate(row_data, 1):
            _write_data_cell(ws, row, col, val, color=clr, bold=bold, bg=bg, mono=mono)

    last_data_row = len(report.rx_results) + 2
    ws.auto_filter.ref = f"A2:L{last_data_row}"
    ws.freeze_panes = "A3"
    # col widths: Band  Blk  Freq  MCS  BW   Ant   OrigRSSI  DUT RSSI  ΔRSSI  OrigPER  DutPER  Status
    _set_col_widths(ws, [10,   5,   10,   9,   8,    14,   13,       13,       10,    11,      11,     12])


def _build_corrections_sheet(ws, report):
    """
    EEPROM Corrections sheet — amber-accented, same 10-column layout as PDF.
    """
    ws.sheet_view.showGridLines = False

    from core.calibration import CalibrationEngine
    corrections = CalibrationEngine.get_corrections(report)

    title_text = "EEPROM Corrections Required" if corrections else "No EEPROM corrections required"
    title_clr  = WARN_CLR if corrections else PASS_CLR
    ws.merge_cells("A1:J1")
    ws.row_dimensions[1].height = 24
    t = ws.cell(row=1, column=1, value=title_text)
    t.font = Font(bold=True, size=12, color=title_clr, name="Arial")
    t.alignment = _align(h="left", v="center")
    t.fill = _fill(ROW_ALT)

    if not corrections:
        msg = ws.cell(row=3, column=1,
                      value="All TX blocks are within tolerance — no EEPROM updates needed.")
        msg.font = _font(bold=True, color=PASS_CLR, size=10)
        _set_col_widths(ws, [14] * 10)
        return

    headers = [
        "BAND", "BLOCK", "FREQ (MHz)", "MOD", "BW", "ANT",
        "ORIGIN dBm", "DUT dBm", "DELTA", "CORRECTION TO APPLY",
    ]
    _write_header_row(ws, 2, headers, bg=HEADER_BG)

    warn_side = Side(border_style="thin", color=WARN_CLR)
    warn_border = Border(left=warn_side, right=warn_side,
                         top=warn_side,  bottom=warn_side)

    for i, c_data in enumerate(corrections):
        row = i + 3
        bg  = WARN_ROW2 if (i % 2 == 1) else WARN_ROW1
        ws.row_dimensions[row].height = 17

        vals = [
            (c_data.get("band", ""),                                  DARK,    False, True),
            (str(c_data.get("label", "")).split(".")[0],              DARK,    False, True),
            (c_data.get("freq_mhz", ""),                              DARK,    False, True),
            (c_data.get("modulation", ""),                            DARK,    False, True),
            (c_data.get("bandwidth", ""),                             DARK,    False, True),
            (c_data.get("antenna", ""),                               DARK,    False, True),
            (f"{c_data.get('origin_measured_dbm', 0):.2f}",          DARK,    False, True),
            (f"{c_data.get('dut_measured_dbm',    0):.2f}",          DARK,    False, True),
            (f"{c_data.get('delta_dbm',           0):+.3f}",         DARK,    False, True),
            (f"{c_data.get('correction_dbm',      0):+.3f}",         WARN_CLR, True, False),
        ]
        for col, (val, clr, bold, mono) in enumerate(vals, 1):
            c = ws.cell(row=row, column=col, value=val)
            c.font = Font(
                bold=bold, color=clr, size=8,
                name="Arial" if bold else "Courier New"
            )
            c.fill = _fill(bg)
            c.alignment = _align()
            c.border = warn_border

    last_data_row = len(corrections) + 2
    ws.auto_filter.ref = f"A2:J{last_data_row}"
    ws.freeze_panes = "A3"
    # col widths: Band  Block  Freq  Mod  BW   Ant   Orig   DUT   Delta  Correction
    _set_col_widths(ws, [10,   7,    10,   10,  8,    8,    13,   13,    11,    20])